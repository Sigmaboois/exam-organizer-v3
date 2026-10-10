"""
Teacher outputs: a marked copy of the student's PDF, and a spreadsheet of results.

The marked PDF gets a cover page (score, grade, every question with its marks and
feedback) and a small red marks box on each page that has answers. Everything is
drawn as real PDF text with a standard font, so it looks sharp in any PDF viewer.
"""

import os
import textwrap

from pypdf import PdfReader, PdfWriter, PageObject
from pypdf.generic import (ContentStream, DecodedStreamObject, DictionaryObject, NameObject)

RED = (0.80, 0.10, 0.10)
DARK = (0.15, 0.15, 0.20)
GREY = (0.40, 0.40, 0.45)
GREEN = (0.10, 0.50, 0.20)
AMBER = (0.75, 0.45, 0.00)


# --------------------------------------------------------------------------
# Tiny PDF text drawing helpers (Helvetica is built into every PDF viewer)
# --------------------------------------------------------------------------

def _pdf_text(value):
    """Make text safe for a PDF string in the standard Helvetica encoding."""
    value = str(value).replace("–", "-").replace("—", "-").replace("’", "'")
    value = value.encode("latin-1", "replace").decode("latin-1")
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


class _Canvas:
    def __init__(self):
        self.ops = []

    def text(self, x, y, value, size=10, bold=False, colour=DARK):
        font = "/F2" if bold else "/F1"
        self.ops.append(f"BT {colour[0]} {colour[1]} {colour[2]} rg {font} {size} Tf {x:.1f} {y:.1f} Td "
                        f"({_pdf_text(value)}) Tj ET")

    def box(self, x, y, width, height, fill=None, stroke=None, line=1.0):
        parts = ["q"]
        if fill:
            parts.append(f"{fill[0]} {fill[1]} {fill[2]} rg")
        if stroke:
            parts.append(f"{stroke[0]} {stroke[1]} {stroke[2]} RG {line} w")
        parts.append(f"{x:.1f} {y:.1f} {width:.1f} {height:.1f} re")
        parts.append("B" if fill and stroke else "f" if fill else "S")
        parts.append("Q")
        self.ops.append(" ".join(parts))

    def page(self, width, height):
        page = PageObject.create_blank_page(width=width, height=height)
        fonts = DictionaryObject()
        for key, base in (("/F1", "/Helvetica"), ("/F2", "/Helvetica-Bold")):
            font = DictionaryObject()
            font[NameObject("/Type")] = NameObject("/Font")
            font[NameObject("/Subtype")] = NameObject("/Type1")
            font[NameObject("/BaseFont")] = NameObject(base)
            font[NameObject("/Encoding")] = NameObject("/WinAnsiEncoding")
            fonts[NameObject(key)] = font
        resources = DictionaryObject()
        resources[NameObject("/Font")] = fonts
        page[NameObject("/Resources")] = resources
        stream = DecodedStreamObject()
        stream.set_data("\n".join(self.ops).encode("latin-1"))
        page[NameObject("/Contents")] = stream
        return page


def _wrap(value, width):
    return textwrap.wrap(str(value or ""), width=width) or [""]


# --------------------------------------------------------------------------
# Marked PDF
# --------------------------------------------------------------------------

def _cover_pages(width, height, info, marking, grade):
    """Cover page(s) with the summary and every question. Returns a list of pages."""
    pages = []
    canvas = _Canvas()
    margin = 50
    y = height - margin

    def new_page():
        nonlocal canvas, y
        pages.append(canvas.page(width, height))
        canvas = _Canvas()
        y = height - margin

    canvas.text(margin, y - 10, "Marked by Exam Organizer", size=20, bold=True)
    y -= 34
    canvas.text(margin, y, info["title"], size=11, colour=GREY)
    y -= 16
    if info.get("student"):
        canvas.text(margin, y, f"Student: {info['student']}", size=11, colour=GREY)
        y -= 16

    percent = round(100 * marking["total"] / marking["max_total"]) if marking["max_total"] else 0
    canvas.box(margin, y - 62, width - 2 * margin, 52, fill=(0.95, 0.96, 1.0), stroke=(0.80, 0.84, 0.95))
    canvas.text(margin + 16, y - 42, f"{marking['total']} / {marking['max_total']}", size=26, bold=True, colour=RED)
    canvas.text(margin + 190, y - 33, f"{percent}%", size=14, bold=True)
    if grade:
        canvas.text(margin + 190, y - 50, f"Grade {grade} (Cambridge thresholds)", size=10, colour=GREY)
    y -= 84

    for line in _wrap(marking.get("summary"), 95):
        canvas.text(margin, y, line, size=10)
        y -= 14
    if marking.get("needs_review"):
        y -= 4
        canvas.text(margin, y, "Please double-check: " + ", ".join(marking["needs_review"]), size=10,
                    bold=True, colour=AMBER)
        y -= 14
    y -= 12

    def table_header():
        nonlocal y
        canvas.text(margin, y, "Question", size=9, bold=True, colour=GREY)
        canvas.text(margin + 70, y, "Marks", size=9, bold=True, colour=GREY)
        canvas.text(margin + 125, y, "Marking and feedback", size=9, bold=True, colour=GREY)
        y -= 14

    table_header()
    for q in marking["questions"]:
        lines = _wrap(q.get("marking"), 80) + _wrap(q.get("feedback"), 80)
        needed = 14 * len(lines) + 10
        if y - needed < margin:
            new_page()
            table_header()
        colour = GREEN if q["marks_awarded"] == q["max_marks"] else RED if q["marks_awarded"] == 0 else AMBER
        canvas.text(margin, y, q["question"], size=10, bold=True)
        canvas.text(margin + 70, y, f"{q['marks_awarded']} / {q['max_marks']}", size=10, bold=True, colour=colour)
        for index, line in enumerate(lines):
            canvas.text(margin + 125, y - 14 * index, line, size=9,
                        colour=DARK if index < len(_wrap(q.get("marking"), 80)) else GREY)
        y -= needed
    pages.append(canvas.page(width, height))
    return pages


def _marks_overlay(width, height, questions):
    """A red box in the top-right corner listing the marks for questions on this page."""
    canvas = _Canvas()
    box_width = 118
    box_height = 18 + 13 * len(questions)
    x = width - box_width - 14
    y = height - box_height - 14
    canvas.box(x, y, box_width, box_height, fill=(1, 0.97, 0.97), stroke=RED)
    canvas.text(x + 8, y + box_height - 13, "Marks", size=8, bold=True, colour=RED)
    for index, q in enumerate(questions):
        canvas.text(x + 8, y + box_height - 27 - 13 * index,
                    f"{q['question']}:  {q['marks_awarded']}/{q['max_marks']}", size=9, bold=True, colour=RED)
    return canvas.page(width, height)


def marked_pdf(student_pdf, marking, output_path, title, student=None, grade=None):
    """Write a marked copy of the student's PDF and return its path."""
    reader = PdfReader(student_pdf)
    writer = PdfWriter()
    first = reader.pages[0]
    width, height = float(first.mediabox.width), float(first.mediabox.height)

    for page in _cover_pages(width, height, {"title": title, "student": student}, marking, grade):
        writer.add_page(page)

    by_page = {}
    for q in marking["questions"]:
        by_page.setdefault(int(q.get("page") or 0), []).append(q)
    for number, page in enumerate(reader.pages, start=1):
        if by_page.get(number):
            page_width, page_height = float(page.mediabox.width), float(page.mediabox.height)
            page.merge_page(_marks_overlay(page_width, page_height, by_page[number]))
        writer.add_page(page)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    with open(output_path, "wb") as file:
        writer.write(file)
    return output_path


# --------------------------------------------------------------------------
# Spreadsheet export
# --------------------------------------------------------------------------

COLUMNS = [("Student", 24), ("Class", 16), ("Qualification", 16), ("Subject", 30), ("Code", 8),
           ("Paper", 8), ("Series", 20), ("Score", 8), ("Out of", 8), ("%", 8), ("Grade", 8), ("Checked on", 18)]


def _sheet(workbook, title, rows):
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = workbook.create_sheet(title=title[:31] or "Results")
    sheet.append([name for name, _ in COLUMNS])
    for row in rows:
        percent = round(100 * row["score"] / row["max_score"], 1) if row["max_score"] else None
        sheet.append([
            row["student_name"], row["class_name"], row["qualification"], row["subject_name"], row["subject_code"],
            f"{row['paper']}{row['variant']}", f"{(row['session'] or '').replace('-', '/')} {row['year'] or ''}".strip(),
            row["score"], row["max_score"], percent, row["grade"], row["checked_at"].replace("T", " "),
        ])
    header_fill = PatternFill("solid", fgColor="2575FC")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def export_spreadsheet(rows, output_path):
    """rows: result dicts from scorebook.results(). One sheet per class plus an 'All results' sheet."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "All results", rows)
    by_class = {}
    for row in rows:
        by_class.setdefault(row["class_name"], []).append(row)
    for class_name in sorted(by_class, key=str.lower):
        safe = "".join(c for c in class_name if c not in '[]:*?/\\')
        _sheet(workbook, safe, by_class[class_name])
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook.save(output_path)
    return output_path
